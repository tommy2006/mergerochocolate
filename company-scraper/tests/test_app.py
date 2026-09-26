"""Offline test of the web app: API, background jobs and exports, with every HTTP
call routed to the saved fixtures (see test_offline.py).

Run:  python tests/test_app.py
"""
import csv
import io
import sys
import tempfile
import time
from pathlib import Path
from unittest import mock

HERE = Path(__file__).parent
sys.path.insert(0, str(HERE.parent))
sys.path.insert(0, str(HERE))
import app as webapp  # noqa: E402
import company_scraper as cs  # noqa: E402
from test_offline import fake_get  # noqa: E402


def wait(client, job_id, timeout=20):
    end = time.time() + timeout
    while time.time() < end:
        job = client.get(f"/api/jobs/{job_id}").get_json()
        if job["status"] != "running":
            return job
        time.sleep(0.05)
    raise AssertionError("job did not finish")


def run():
    with tempfile.TemporaryDirectory() as tmp, \
            mock.patch.object(cs, "http_get", side_effect=fake_get), mock.patch.object(cs.time, "sleep"):
        client = webapp.create_app(Path(tmp)).test_client()

        assert client.get("/").status_code == 200
        assert client.get("/api/ping").get_json() == {"app": "company-scraper"}
        assert client.get("/api/companies").get_json() == []

        # single lookup: progress log, saved profile
        r = client.post("/api/jobs", json={"kind": "lookup", "queries": ["0180611-0"]})
        assert r.status_code == 201
        job = wait(client, r.get_json()["id"])
        item = job["items"][0]
        assert job["status"] == "done" and item["status"] == "found", job
        assert item["business_id"] == "0180611-0"
        assert any("Looking up" in line for line in item["log"])
        assert (Path(tmp) / "0180611-0.json").exists()

        rows = client.get("/api/companies").get_json()
        assert len(rows) == 1 and rows[0]["name"] == "Lapuan Kuljetus Oy"
        assert rows[0]["revenue"] == 1204342.08 and rows[0]["revenue_growth_pct"] == -41.1
        profile = client.get("/api/companies/0180611-0").get_json()
        assert profile["registry"]["municipality"] == "LAPUA"

        # batch: duplicates skipped, unknown company reported, logs left out of polling
        r = client.post("/api/jobs", json={"kind": "batch",
                                           "queries": ["0180611-0", "Lapuan Kuljetus", "Nonexistent Firma Oy"],
                                           "with_website": False})
        job = wait(client, r.get_json()["id"])
        statuses = [i["status"] for i in job["items"]]
        assert statuses == ["found", "duplicate", "not_found"], statuses
        assert job["items"][1]["duplicate_of_row"] == 1
        assert "log" not in job["items"][0]
        assert job["counts"]["found"] == 1
        # the batch skipped the website: the website section saved by the first lookup is kept
        saved = client.get("/api/companies/0180611-0").get_json()
        assert saved["website"]["available"] and saved["website"]["retrieved_at"] <= saved["scraped_at"]
        assert saved["financials"]["retrieved_at"] >= saved["website"]["retrieved_at"]

        # exports
        r = client.get("/api/export?format=csv&ids=0180611-0")
        assert r.status_code == 200 and "attachment" in r.headers["Content-Disposition"]
        table = list(csv.reader(io.StringIO(r.data.decode("utf-8-sig"))))
        assert table[0][:2] == ["Business ID", "Name"] and table[1][:2] == ["0180611-0", "Lapuan Kuljetus Oy"]
        assert len(table) == 2

        r = client.get("/api/export?format=xlsx")
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(r.data))
        ws = wb["Companies"]
        assert ws["A2"].value == "0180611-0" and ws["B2"].value == "Lapuan Kuljetus Oy"
        rev_col = [c.value for c in ws[1]].index("Revenue €") + 1
        assert ws.cell(row=2, column=rev_col).value == 1204342.08

        # Evidence sheet: one row per fact with its source and quote / filed fact / formula
        ev = wb["Evidence"]
        hdr = [c.value for c in ev[1]]
        facts = [dict(zip(hdr, (c.value for c in row))) for row in ev.iter_rows(min_row=2)]
        by_field = {f["Field"]: f for f in facts}
        rev = by_field["Revenue, year ending 2025-12-31"]
        assert rev["Value"] == 1204342.08 and "x673" in rev["Evidence"] and "1204342.08" in rev["Evidence"]
        assert rev["Source URL"].startswith("https://avoindata.prh.fi/opendata-xbrl-api/")
        assert by_field["EBITDA, year ending 2025-12-31"]["Evidence"] == "Calculated: Operating profit + depreciation"
        email = by_field["Email"]
        assert email["Value"] == "info@lapuankuljetus.fi" and email["Source URL"].endswith("/yhteystiedot")
        assert "info@lapuankuljetus.fi" in email["Evidence"]
        assert by_field["Employer register"]["Value"] == "Ended 2023-02-28"
        assert "ended 2023-02-28" in by_field["Employer register"]["Evidence"]
        assert by_field["Name"]["Source URL"] == "https://tietopalvelu.ytj.fi/yritys/0180611-0"
        assert by_field["Website"]["Evidence"].startswith("Guessed from the company name")
        assert ev.cell(row=2, column=hdr.index("Source URL") + 1).hyperlink is not None

        assert client.get("/api/export?format=json").get_json()[0]["registry"]["business_id"] == "0180611-0"

        # scraped text that looks like a formula stays text
        row = webapp.flatten_profile(profile)
        row["about"] = "=HYPERLINK(\"http://evil\")"
        ws = load_workbook(io.BytesIO(webapp.to_xlsx([row]))).active
        about_col = [c.value for c in ws[1]].index("About") + 1
        assert ws.cell(row=2, column=about_col).data_type == "s"
        assert "'=HYPERLINK" in webapp.to_csv([row]).decode("utf-8-sig")

        # bad input
        assert client.post("/api/jobs", json={"queries": []}).status_code == 400
        assert client.get("/api/companies/..%2Fapp").status_code == 404
        assert client.get("/api/companies/1234567-8").status_code == 404
        assert client.get("/api/jobs/nope").status_code == 404

    print("ALL APP CHECKS PASSED")


if __name__ == "__main__":
    run()
