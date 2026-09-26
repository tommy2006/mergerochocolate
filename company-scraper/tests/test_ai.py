"""Offline test of the AI features. Claude's replies are simulated: nothing is sent to Anthropic.

Run:  python tests/test_ai.py
"""
import io
import sys
import tempfile
from pathlib import Path
from unittest import mock

HERE = Path(__file__).parent
sys.path.insert(0, str(HERE.parent))
sys.path.insert(0, str(HERE))
import ai  # noqa: E402
import app as webapp  # noqa: E402
import company_scraper as cs  # noqa: E402
from test_app import wait  # noqa: E402
from test_offline import fake_get  # noqa: E402

seen = {}


def fake_analyze(facts, name, business_id, detail="brief"):
    seen["facts"], seen["name"], seen["detail"] = facts, name, detail
    rev = next(f["id"] for f in facts if f["field"] == "Revenue, year ending 2025-12-31")
    return {"analysis": {
        "headline": f"Revenue fell 41% in 2025 [{rev}].", "summary": f"A Lapua haulier [{rev}].",
        "business": [], "financial_health": {"rating": "weak", "points": [f"Revenue down [{rev}]"]},
        "sale_signals": {"rating": "high", "points": []}, "strengths": [], "risks": [],
        "questions": ["Who owns the company?"], "data_gaps": []},
        "model": "claude-opus-5", "usage": {"input_tokens": 1, "output_tokens": 1}}


def fake_ask(facts, name, business_id, history, question, on_text=None, on_reset=None):
    seen["history"], seen["question"] = history, question
    on_text("declined attempt")  # a fallback hand-over drops text from the declined attempt
    on_reset()
    on_text("Revenue fell [F1]")
    return {"answer": "Revenue fell [F1].", "model": "claude-opus-5", "usage": {}}


def run():
    # citations: groups are split, unknown IDs dropped; spreadsheets get plain text
    assert ai.clean_citations("Up 5% [F1, F2] and [F9].", {"F1", "F2"}) == "Up 5% [F1][F2] and ."
    assert ai.strip_citations("Revenue fell 41% [F1][F2].") == "Revenue fell 41%."

    with tempfile.TemporaryDirectory() as tmp, \
            mock.patch.object(cs, "http_get", side_effect=fake_get), mock.patch.object(cs.time, "sleep"), \
            mock.patch.object(ai, "is_configured", return_value=True), \
            mock.patch.object(ai, "analyze", side_effect=fake_analyze), \
            mock.patch.object(ai, "ask", side_effect=fake_ask):
        client = webapp.create_app(Path(tmp)).test_client()
        assert client.get("/api/ai/status").get_json()["configured"] is True
        wait(client, client.post("/api/jobs", json={"kind": "lookup", "queries": ["0180611-0"]}).get_json()["id"])

        # analysis: Claude gets numbered facts with their evidence; the result is saved with them
        r = client.post("/api/companies/0180611-0/analysis", json={})
        assert r.status_code == 201
        job = wait(client, r.get_json()["id"])
        assert job["items"][0]["status"] == "done", job
        facts = seen["facts"]
        assert seen["name"] == "Lapuan Kuljetus Oy"
        assert [f["id"] for f in facts[:3]] == ["F1", "F2", "F3"] and len({f["id"] for f in facts}) == len(facts)
        block = ai.facts_block(facts)
        assert "<facts>" in block and "info@lapuankuljetus.fi" in block and "evidence:" in block
        saved = client.get("/api/companies/0180611-0").get_json()
        analysis = saved["ai_analysis"]
        assert analysis["based_on"] == saved["scraped_at"] and analysis["facts"] == facts
        assert analysis["analysis"]["financial_health"]["rating"] == "weak"
        assert analysis["detail"] == "brief" and seen["detail"] == "brief"

        # "More detail" asks for the detailed version
        r = client.post("/api/companies/0180611-0/analysis", json={"detail": "detailed"})
        wait(client, r.get_json()["id"])
        assert seen["detail"] == "detailed"
        assert client.get("/api/companies/0180611-0").get_json()["ai_analysis"]["detail"] == "detailed"

        # the analysis shows up in the list and in Excel, without citation markers
        row = client.get("/api/companies").get_json()[0]
        assert row["ai_headline"] == "Revenue fell 41% in 2025." and row["ai_sale_signals"] == "high"
        from openpyxl import load_workbook
        ws = load_workbook(io.BytesIO(client.get("/api/export?format=xlsx").data))["Companies"]
        header = [c.value for c in ws[1]]
        assert ws.cell(row=2, column=header.index("AI headline") + 1).value == "Revenue fell 41% in 2025."

        # a refresh keeps the analysis, which records the data it was based on
        wait(client, client.post("/api/jobs", json={"kind": "lookup", "queries": ["0180611-0"]}).get_json()["id"])
        assert client.get("/api/companies/0180611-0").get_json()["ai_analysis"]["based_on"] == analysis["based_on"]

        # questions: earlier turns are passed on; the answer comes back with the facts it cites
        r = client.post("/api/companies/0180611-0/ask",
                        json={"question": "Why?", "history": [{"q": "What do they do?", "a": "Haulage."}]})
        item = wait(client, r.get_json()["id"])["items"][0]
        assert item["status"] == "done" and item["answer"] == "Revenue fell [F1]." and set(item["facts"]) == {"F1"}
        assert seen["history"] == [{"q": "What do they do?", "a": "Haulage."}] and seen["question"] == "Why?"
        assert client.post("/api/companies/0180611-0/ask", json={"question": " "}).status_code == 400

        # batch with AI: each company found is analysed
        r = client.post("/api/jobs", json={"kind": "batch", "queries": ["0180611-0", "Lapuan Kuljetus"],
                                           "with_ai": True, "with_website": False})
        job = wait(client, r.get_json()["id"])
        assert job["items"][0]["row"]["ai_headline"] == "Revenue fell 41% in 2025."
        assert job["items"][1]["status"] == "duplicate"

        # a failed call is reported, not raised
        with mock.patch.object(ai, "analyze", side_effect=ai.AIError("The Anthropic API key was rejected.")):
            item = wait(client, client.post("/api/companies/0180611-0/analysis", json={}).get_json()["id"])["items"][0]
        assert item["status"] == "error" and "rejected" in item["note"]

        # other websites and other host names can't start scrapes or paid AI calls
        assert client.post("/api/companies/0180611-0/analysis", data="{}", content_type="text/plain").status_code == 403
        assert client.post("/api/companies/0180611-0/analysis", json={},
                           headers={"Origin": "https://evil.example"}).status_code == 403
        assert client.post("/api/companies/0180611-0/analysis", json={},
                           headers={"Origin": "http://localhost"}).status_code == 201
        assert client.get("/api/ping", headers={"Host": "evil.example:8765"}).status_code == 403
        assert client.post("/api/companies/9999999-9/analysis", json={}).status_code == 404

    with mock.patch.object(ai, "is_configured", return_value=False), tempfile.TemporaryDirectory() as tmp:
        client = webapp.create_app(Path(tmp)).test_client()
        assert client.get("/api/ai/status").get_json()["configured"] is False

    print("ALL AI CHECKS PASSED")


if __name__ == "__main__":
    run()
