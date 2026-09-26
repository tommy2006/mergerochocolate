#!/usr/bin/env python3
"""
Company Scraper – a local web app around company_scraper.py.

Start:  python app.py                 (opens http://127.0.0.1:8765 in your browser)
        python app.py --port 9000 --no-browser

Look up one company or run a batch, browse saved profiles and download them as
Excel, CSV or JSON. Profiles are saved to output/<BusinessID>.json – the same
files the command line writes. With an Anthropic API key (see ai.py), Claude writes
a cited analysis of each company and answers questions about it.
"""

from __future__ import annotations

import argparse
import copy
import csv
import io
import json
import logging
import socket
import sys
import threading
import time
import traceback
import uuid
import webbrowser
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime
from pathlib import Path
from urllib.parse import urlparse

import requests
from flask import Flask, Response, abort, jsonify, request, send_from_directory
from werkzeug.serving import make_server

import ai
import buyers
import company_scraper as cs

HERE = Path(__file__).resolve().parent
STATIC = HERE / "static"
APP_ID = "company-scraper"
DEFAULT_PORT = 8765
MAX_BATCH = 500
MAX_JOBS_KEPT = 50
KEEP_LOG = {"lookup", "find", "add-buyer"}  # job kinds whose progress messages the page shows
SAVE_LOCK = threading.Lock()  # one read-modify-write of a profile file at a time


# --------------------------------------------------------------------------
# Flat rows (sidebar list, Excel / CSV export)
# --------------------------------------------------------------------------

# (key, column title, kind) – kind picks the Excel number format
EXPORT_COLUMNS = [
    ("business_id", "Business ID", "text"),
    ("name", "Name", "text"),
    ("company_form", "Company form", "text"),
    ("industry_code", "Industry code", "text"),
    ("industry", "Industry", "text"),
    ("municipality", "Municipality", "text"),
    ("street_address", "Street address", "text"),
    ("postal_address", "Postal address", "text"),
    ("registered_on", "Registered on", "text"),
    ("company_age_years", "Age (years)", "num1"),
    ("active", "Active", "bool"),
    ("status", "Status", "text"),
    ("in_employer_register", "Employer register", "bool"),
    ("vat_registered", "VAT register", "bool"),
    ("website", "Website", "text"),
    ("fiscal_year_end", "Financial year end", "text"),
    ("revenue", "Revenue €", "money"),
    ("revenue_prev", "Revenue previous year €", "money"),
    ("revenue_growth_pct", "Revenue growth %", "num1"),
    ("ebitda", "EBITDA €", "money"),
    ("ebitda_margin_pct", "EBITDA margin %", "num1"),
    ("operating_profit", "Operating profit €", "money"),
    ("net_profit", "Net profit €", "money"),
    ("personnel_costs", "Personnel costs €", "money"),
    ("total_assets", "Total assets €", "money"),
    ("equity", "Equity €", "money"),
    ("equity_ratio_pct", "Equity ratio %", "num1"),
    ("emails", "Emails", "text"),
    ("phones", "Phones", "text"),
    ("social_links", "Social links", "text"),
    ("people", "People mentioned", "text"),
    ("signals", "Signals", "text"),
    ("about", "About", "text"),
    ("ai_headline", "AI headline", "text"),
    ("ai_financial_health", "AI financial health", "text"),
    ("ai_sale_signals", "AI sale signals", "text"),
    ("ai_summary", "AI summary", "text"),
    ("scraped_at", "Scraped at", "text"),
]

SIGNAL_LABELS = {
    "family_business": "Family business",
    "multi_generation": "Multi-generation",
    "hiring": "Hiring",
    "growth_or_expansion": "Growth / expansion",
    "succession_or_sale": "Succession / sale",
    "certifications": "Certifications",
}


def flatten_profile(p: dict) -> dict:
    """One flat row per company, latest financial year first."""
    r = p["registry"]
    fin = p.get("financials") or {}
    years = fin.get("years") or {}
    keys = sorted(years, reverse=True)
    cur = years[keys[0]] if keys else {}
    prev = years[keys[1]] if len(keys) > 1 else {}
    web = p.get("website") or {}
    web_ok = web.get("available")
    analysis = (p.get("ai_analysis") or {}).get("analysis") or {}

    status = ["Active" if r.get("active") else "Not active"]
    status += [f"{s['type']} since {s['since']}" for s in r.get("situations", []) if not s.get("until")]
    if r.get("trade_register_status"):
        status.append(r["trade_register_status"])

    return {
        "business_id": r["business_id"],
        "name": r.get("name"),
        "company_form": r.get("company_form"),
        "industry_code": r.get("industry_code"),
        "industry": r.get("industry"),
        "municipality": r.get("municipality"),
        "street_address": r.get("street_address"),
        "postal_address": r.get("postal_address"),
        "registered_on": r.get("registered_on"),
        "company_age_years": r.get("company_age_years"),
        "active": r.get("active"),
        "status": " · ".join(status),
        "in_employer_register": r.get("in_employer_register"),
        "vat_registered": r.get("vat_registered"),
        "website": (web.get("url") if web_ok else None) or r.get("website"),
        "fiscal_year_end": keys[0] if keys else None,
        "revenue": cur.get("revenue"),
        "revenue_prev": prev.get("revenue"),
        "revenue_growth_pct": cur.get("revenue_growth_pct"),
        "ebitda": cur.get("ebitda"),
        "ebitda_margin_pct": cur.get("ebitda_margin_pct"),
        "operating_profit": cur.get("operating_profit"),
        "net_profit": cur.get("net_profit"),
        "personnel_costs": cur.get("personnel_costs"),
        "total_assets": cur.get("total_assets"),
        "equity": cur.get("equity"),
        "equity_ratio_pct": cur.get("equity_ratio_pct"),
        "emails": "; ".join(web.get("emails") or []) if web_ok else None,
        "phones": "; ".join(web.get("phones") or []) if web_ok else None,
        "social_links": "; ".join(web.get("social_links") or []) if web_ok else None,
        "people": " | ".join(web.get("people_mentions") or []) if web_ok else None,
        "signals": ", ".join(SIGNAL_LABELS.get(k, k) for k in (web.get("signals") or {})) if web_ok else None,
        "about": (web.get("meta_description") or web.get("about_text")) if web_ok else None,
        "ai_headline": ai.strip_citations(analysis["headline"]) if analysis.get("headline") else None,
        "ai_financial_health": (analysis.get("financial_health") or {}).get("rating"),
        "ai_sale_signals": (analysis.get("sale_signals") or {}).get("rating"),
        "ai_summary": ai.strip_citations(analysis["summary"]) if analysis.get("summary") else None,
        "scraped_at": p.get("scraped_at"),
    }


FIN_LABELS = {
    "revenue": "Revenue", "revenue_growth_pct": "Revenue growth %",
    "other_operating_income": "Other operating income", "materials_and_services": "Materials and services",
    "personnel_costs": "Personnel costs", "wages_and_salaries": "Wages and salaries",
    "depreciation": "Depreciation", "other_operating_expenses": "Other operating expenses",
    "ebitda": "EBITDA", "ebitda_margin_pct": "EBITDA margin %", "operating_profit": "Operating profit",
    "operating_margin_pct": "Operating margin %", "profit_before_taxes": "Profit before taxes",
    "income_taxes": "Income taxes", "net_profit": "Net profit", "net_margin_pct": "Net margin %",
    "total_assets": "Total assets", "equity": "Equity", "equity_ratio_pct": "Equity ratio %",
    "total_liabilities": "Total liabilities",
}

EVIDENCE_COLUMNS = [("business_id", "Business ID"), ("company", "Company"), ("section", "Section"),
                    ("field", "Field"), ("value", "Value"), ("source", "Source"), ("url", "Source URL"),
                    ("evidence", "Evidence"), ("retrieved_at", "Retrieved at")]


def evidence_rows(p: dict) -> list[dict]:
    """One row per stated fact: where it came from and the quote, filed fact or formula behind it."""
    r = p["registry"]
    rows = []
    fin = p.get("financials") or {}
    web = p.get("website") or {}
    when = {"Company": p.get("scraped_at"), "Financials": fin.get("retrieved_at") or p.get("scraped_at"),
            "Website": web.get("retrieved_at") or p.get("scraped_at")}

    def add(section, field, value, source, url, evidence):
        if value in (None, "", []):
            return
        rows.append({"business_id": r["business_id"], "company": r.get("name"), "section": section,
                     "field": field, "value": value, "source": source, "url": url, "evidence": evidence,
                     "retrieved_at": when[section]})

    reg, reg_url = "PRH trade register (YTJ)", r.get("register_page") or f"{cs.YTJ_PAGE}/{r['business_id']}"
    record = (f"Official PRH record, last changed {r['registry_last_modified'][:10]}"
              if r.get("registry_last_modified") else "Official PRH record")
    for field, value in [
        ("Name", r.get("name")), ("Company form", r.get("company_form")),
        ("Industry", f"{r['industry']} ({r['industry_code']})" if r.get("industry") else None),
        ("Street address", r.get("street_address")), ("Postal address", r.get("postal_address")),
        ("Registered on", r.get("registered_on")), ("Business ID granted on", r.get("business_id_granted_on")),
        ("Trade register status", r.get("trade_register_status")), ("Website in registry", r.get("website")),
    ]:
        add("Company", field, value, reg, reg_url, record)
    for sit in r.get("situations", []):
        add("Company", "Situation", sit["type"], reg, reg_url,
            f"Registered {sit['since']}" + (f", ended {sit['until']}" if sit.get("until") else ", still valid"))
    for e in r.get("registers", []):
        ended = e.get("until")
        add("Company", e["register"], f"Ended {ended}" if ended else (e.get("status") or "Registered"), reg, reg_url,
            f"Entry from {e['since']}" + (f", ended {ended}" if ended else ", still valid"))

    if fin.get("available"):
        facts, calc = fin.get("evidence") or {}, fin.get("calculated") or cs.CALCULATED
        for year, items in fin["years"].items():
            for item, value in items.items():
                if value is None:
                    continue
                fact = (facts.get(year) or {}).get(item)
                if fact:
                    text = (f"Filed fact {fact['element']} in context {fact['context']} "
                            f"(line item {fact['code']}), filed value {fact['filed_value']}")
                elif item in calc:
                    text = f"Calculated: {calc[item]}"
                else:
                    text = "Refresh the profile to record the filed fact"
                add("Financials", f"{FIN_LABELS.get(item, item)}, year ending {year}", value,
                    "PRH digital financial statement (XBRL)", fin.get("source"), text)

    if web.get("available"):
        wev, site = web.get("evidence") or {}, "Company website"
        how = {"registry": "Listed as the website in the PRH record",
               "given": "Website entered by the user"}.get(web.get("url_source"))
        if not how and web.get("url_evidence"):
            ue = web["url_evidence"]
            how = (f"Guessed from the company name; accepted because the page mentions the {ue['verified_by']}"
                   + (f": {ue['snippet']}" if ue.get("snippet") else ""))
        add("Website", "Website", web.get("url"), site, web.get("url"), how or web.get("url_source"))
        about = web.get("meta_description") or web.get("about_text")
        if about:
            if web.get("meta_description"):
                where = "Meta description of the home page"
            else:
                where = "Paragraphs on " + ", ".join((wev.get("about_text") or {}).get("pages") or [web["url"]])
            add("Website", "About", about, site, web.get("url"), where)
        for kind, field in [("emails", "Email"), ("phones", "Phone"), ("social_links", "Social link"),
                            ("people_mentions", "Person / role"), ("founding_years_mentioned", "Founding year")]:
            for v in web.get(kind) or []:
                e = (wev.get(kind) or {}).get(v) or {}
                add("Website", field, v, site, e.get("page") or web.get("url"),
                    e.get("snippet") or (f"Found in the {e['found_in']}" if e.get("found_in") else ""))
        for key, snippet in (web.get("signals") or {}).items():
            e = (wev.get("signals") or {}).get(key) or {}
            add("Website", "Signal", SIGNAL_LABELS.get(key, key), site, e.get("page") or web.get("url"),
                e.get("snippet") or snippet)
    return rows


def _cell(value, kind: str):
    if value is None or value == "":
        return None
    if kind == "bool":
        return "Yes" if value else "No"
    if kind == "money":
        return round(value, 2)
    return value


def to_xlsx(rows: list[dict], evidence: list[dict] | None = None) -> bytes:
    from openpyxl import Workbook
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    def put(ws, i, j, value):
        if isinstance(value, str):
            value = ILLEGAL_CHARACTERS_RE.sub("", value)[:32000]
        c = ws.cell(row=i, column=j, value=value)
        if isinstance(value, str) and value.startswith("="):
            c.data_type = "s"  # scraped text, never a formula
        return c

    wb = Workbook()
    ws = wb.active
    ws.title = "Companies"
    ws.append([title for _, title, _ in EXPORT_COLUMNS])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for i, row in enumerate(rows, start=2):
        for j, (key, _, kind) in enumerate(EXPORT_COLUMNS, start=1):
            c = put(ws, i, j, _cell(row.get(key), kind))
            if kind == "money":
                c.number_format = "#,##0"
            elif kind == "num1":
                c.number_format = "0.0"
    widths = {"name": 32, "industry": 40, "street_address": 32, "postal_address": 32, "status": 28,
              "website": 30, "emails": 30, "phones": 24, "social_links": 30, "people": 40,
              "signals": 30, "about": 60, "ai_headline": 60, "ai_summary": 80}
    for j, (key, title, _) in enumerate(EXPORT_COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(j)].width = widths.get(key, max(12, len(title) + 2))
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions

    if evidence is not None:  # one row per fact: where it came from
        ev = wb.create_sheet("Evidence")
        ev.append([title for _, title in EVIDENCE_COLUMNS])
        for cell in ev[1]:
            cell.font = Font(bold=True)
        for i, row in enumerate(evidence, start=2):
            for j, (key, _) in enumerate(EVIDENCE_COLUMNS, start=1):
                c = put(ev, i, j, row.get(key))
                if key == "url" and isinstance(row.get(key), str) and row[key].startswith(("http://", "https://")):
                    c.hyperlink = row[key]
                    c.style = "Hyperlink"
        for j, width in enumerate([12, 26, 11, 30, 30, 26, 40, 90, 20], start=1):
            ev.column_dimensions[get_column_letter(j)].width = width
        ev.freeze_panes = "E2"
        ev.auto_filter.ref = ev.dimensions
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_csv(rows: list[dict]) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([title for _, title, _ in EXPORT_COLUMNS])
    for row in rows:
        out = []
        for key, _, kind in EXPORT_COLUMNS:
            v = _cell(row.get(key), kind)
            if isinstance(v, str) and v[:1] in ("=", "+", "-", "@"):
                v = "'" + v  # stop spreadsheets from running scraped text as a formula
            out.append("" if v is None else v)
        w.writerow(out)
    return ("﻿" + buf.getvalue()).encode("utf-8")  # BOM: Excel reads ä/ö correctly


# --------------------------------------------------------------------------
# Background jobs (one thread per lookup or batch)
# --------------------------------------------------------------------------

class Job:
    def __init__(self, kind: str, queries: list[str], options: dict, meta: dict | None = None):
        self.id = uuid.uuid4().hex[:12]
        self.kind = kind
        self.options = options
        self.meta = meta or {}  # {"title", "bid", "name"} for the running-tasks list
        self.items = [{"row_no": n, "query": q, "status": "queued", "log": []}
                      for n, q in enumerate(queries, 1)]
        self.status = "running"
        self.cancel = threading.Event()
        self.started = time.time()
        self.finished: float | None = None
        self.lock = threading.Lock()

    def update(self, item: dict, **fields) -> None:
        with self.lock:
            item.update(fields)

    def add_log(self, item: dict, msg: str) -> None:
        print(msg, file=sys.stderr)
        with self.lock:
            item["log"] = (item["log"] + [msg])[-200:]

    def finish(self) -> None:
        with self.lock:
            self.status = "stopped" if self.cancel.is_set() else "done"
            self.finished = time.time()

    def to_json(self) -> dict:
        with self.lock:
            items = copy.deepcopy(self.items)
            status, finished = self.status, self.finished
        counts = {s: 0 for s in ("queued", "running", "found", "not_found", "duplicate", "error", "skipped")}
        for it in items:
            counts[it["status"]] = counts.get(it["status"], 0) + 1
            steps = [l.strip()[1:].strip() for l in it["log"] if l.strip().startswith("→")]
            it["step"] = steps[-1] if steps else None
            if self.kind not in KEEP_LOG:
                del it["log"]  # keep batch polling light
        return {
            "id": self.id,
            "kind": self.kind,
            "status": status,
            "elapsed": round((finished or time.time()) - self.started, 1),
            "counts": counts,
            "items": items,
            "meta": self.meta,
        }


# --------------------------------------------------------------------------
# AI (Claude)
# --------------------------------------------------------------------------

def ai_facts(p: dict) -> list[dict]:
    """The facts Claude may use, numbered F1, F2 …: the Evidence rows plus what is missing."""
    rows = evidence_rows(p)
    fin, web = p.get("financials") or {}, p.get("website") or {}
    extra = []
    if fin and not fin.get("available"):
        extra.append(("Financials", "Digital financial statements", f"Not available: {fin.get('note') or 'none found'}",
                      "PRH digital financial statements", fin.get("source"), fin.get("retrieved_at")))
    if web and not web.get("available"):
        extra.append(("Website", "Company website", f"Not available: {web.get('note') or 'none found'}",
                      "Company website", web.get("url"), web.get("retrieved_at")))
    if web.get("available") and web.get("headings"):
        extra.append(("Website", "Page headings", " | ".join(web["headings"]), "Company website", web.get("url"),
                      web.get("retrieved_at")))
    rows += [{"section": sec, "field": field, "value": value, "source": src, "url": url, "evidence": "",
              "retrieved_at": when or p.get("scraped_at")} for sec, field, value, src, url, when in extra]
    keys = ("section", "field", "value", "source", "url", "evidence", "retrieved_at")
    return [{"id": f"F{n}", **{k: row.get(k) for k in keys}} for n, row in enumerate(rows, 1)]


def analyze_profile(p: dict, out_dir: Path, detail: str = "brief") -> dict:
    """Have Claude analyse a profile; the analysis and the facts it cites are saved in the profile."""
    facts = ai_facts(p)
    r = p["registry"]
    result = ai.analyze(facts, r.get("name") or r["business_id"], r["business_id"], detail)
    record = {**result, "detail": detail, "created_at": datetime.now().isoformat(timespec="seconds"),
              "based_on": p.get("scraped_at"), "facts": facts}
    p["ai_analysis"] = record
    with SAVE_LOCK:  # store into the current file, which a refresh may have replaced meanwhile
        path = Path(out_dir) / f"{r['business_id']}.json"
        current = json.loads(path.read_text(encoding="utf-8")) if path.exists() else p
        current["ai_analysis"] = record
        path.write_text(json.dumps(current, ensure_ascii=False, indent=2), encoding="utf-8")
    return record


def _load_profile(out_dir: Path, bid: str) -> dict:
    return json.loads((Path(out_dir) / f"{bid}.json").read_text(encoding="utf-8"))


def run_analysis_job(job: Job, out_dir: Path) -> None:
    item = job.items[0]
    job.update(item, status="running")
    try:
        analyze_profile(_load_profile(out_dir, item["query"]), out_dir, job.options.get("detail", "brief"))
        job.update(item, status="done")
    except ai.AIError as exc:
        job.update(item, status="error", note=str(exc))
    except Exception as exc:
        traceback.print_exc()
        job.update(item, status="error", note=f"Unexpected error: {exc.__class__.__name__}: {exc}")
    job.finish()


def run_ask_job(job: Job, out_dir: Path) -> None:
    """Answer one question; the answer grows in item["answer"] while Claude writes it."""
    item = job.items[0]
    job.update(item, status="running", answer="")

    def add(text: str) -> None:
        with job.lock:
            item["answer"] += text

    def reset() -> None:
        with job.lock:
            item["answer"] = ""

    try:
        p = _load_profile(out_dir, job.options["business_id"])
        facts = ai_facts(p)
        r = p["registry"]
        result = ai.ask(facts, r.get("name") or r["business_id"], r["business_id"],
                        job.options["history"], item["query"], add, reset)
        cited = {f["id"]: f for f in facts if f"[{f['id']}]" in result["answer"]}
        job.update(item, status="done", answer=result["answer"], facts=cited, model=result["model"])
    except ai.AIError as exc:
        job.update(item, status="error", note=str(exc))
    except Exception as exc:
        traceback.print_exc()
        job.update(item, status="error", note=f"Unexpected error: {exc.__class__.__name__}: {exc}")
    job.finish()


def _run_safely(job: Job, item: dict, work) -> None:
    """Run work(); report AI, register and unexpected errors on the job item."""
    try:
        work()
    except ai.AIError as exc:
        job.update(item, status="error", note=str(exc))
    except RuntimeError as exc:  # PRH lookup failed
        job.update(item, status="error", note=str(exc))
    except Exception as exc:
        traceback.print_exc()
        job.update(item, status="error", note=f"Unexpected error: {exc.__class__.__name__}: {exc}")


def make_brief(p: dict, out_dir: Path) -> dict:
    """Claude's seller story and buyer search plan, saved in the campaign."""
    facts = ai_facts(p)
    r = p["registry"]
    result = ai.seller_brief(facts, r.get("name") or r["business_id"], r["business_id"])
    record = {**result["brief"], "model": result["model"], "usage": result["usage"], "facts": facts,
              "created_at": datetime.now().isoformat(timespec="seconds"), "based_on": p.get("scraped_at")}

    def save(c):
        c["brief"] = record
        c["teaser"] = record["teaser"]
    buyers.update_campaign(out_dir, r["business_id"], save)
    return record


def run_brief_job(job: Job, out_dir: Path) -> None:
    item = job.items[0]
    job.update(item, status="running")

    def work():
        make_brief(_load_profile(out_dir, item["query"]), out_dir)
        job.update(item, status="done")
    _run_safely(job, item, work)
    job.finish()


def run_find_job(job: Job, out_dir: Path) -> None:
    """Search the register for buyers, then let Claude score the shortlist."""
    item = job.items[0]
    bid = item["query"]
    job.update(item, status="running")
    token = cs.LOG_HANDLER.set(lambda msg: job.add_log(item, msg))

    def work():
        p = _load_profile(out_dir, bid)
        brief = buyers.load_campaign(out_dir, bid).get("brief")
        if not brief:
            cs.log("→ Claude is writing the seller story and the buyer search plan…")
            brief = make_brief(p, out_dir)
        found, stats = buyers.find_candidates(p, brief, cs.log)
        if found:
            cs.log(f"→ Claude is scoring {len(found)} candidates…")
            scored = ai.score_buyers(buyers.seller_text(p, brief), buyers.buyers_text(found))
            found = buyers.apply_scores(found, scored["buyers"])
            buyers.auto_select(found)

        def save(c):
            manual = [b for b in c.get("buyers") or [] if b.get("added") == "manual"
                      and b["business_id"] not in {f["business_id"] for f in found}]
            c["buyers"] = sorted(found + manual, key=lambda b: b.get("fit") or 0, reverse=True)
            c["search"] = stats
        buyers.update_campaign(out_dir, bid, save)
        job.update(item, status="done", note=None if found else "No matching companies were found.")
    try:
        _run_safely(job, item, work)
    finally:
        cs.LOG_HANDLER.reset(token)
    job.finish()


def run_add_buyer_job(job: Job, out_dir: Path) -> None:
    item = job.items[0]
    bid = job.options["seller"]
    job.update(item, status="running")
    token = cs.LOG_HANDLER.set(lambda msg: job.add_log(item, msg))

    def work():
        cs.log(f"→ Looking up '{item['query']}' and reading its website…")
        record = buyers.enrich_one(item["query"])
        if record is None:
            job.update(item, status="error", note="No company found in the PRH register.")
            return
        if record["business_id"] == bid:
            job.update(item, status="error", note="That's the company being sold.")
            return
        p = _load_profile(out_dir, bid)
        brief = buyers.load_campaign(out_dir, bid).get("brief")
        cs.log("→ Claude is scoring it…")
        scored = ai.score_buyers(buyers.seller_text(p, brief), buyers.buyers_text([record]))
        record = buyers.apply_scores([record], scored["buyers"])[0]
        record["selected"] = True

        def save(c):
            c["buyers"] = sorted([b for b in c.get("buyers") or [] if b["business_id"] != record["business_id"]]
                                 + [record], key=lambda b: b.get("fit") or 0, reverse=True)
        buyers.update_campaign(out_dir, bid, save)
        job.update(item, status="done", name=record["name"])
    try:
        _run_safely(job, item, work)
    finally:
        cs.LOG_HANDLER.reset(token)
    job.finish()


def run_emails_job(job: Job, out_dir: Path) -> None:
    """Write one email per buyer (a few at a time)."""
    bid = job.options["seller"]
    p = _load_profile(out_dir, bid)
    c = buyers.load_campaign(out_dir, bid)
    by_id = {b["business_id"]: b for b in c.get("buyers") or []}
    language, anonymous, sender = c.get("language", "en"), c.get("anonymous", True), c.get("sender") or {}
    seller = buyers.seller_text(p, c.get("brief"), c.get("teaser") or "")

    def one(item):
        buyer = by_id.get(item["query"])
        job.update(item, status="running", name=(buyer or {}).get("name"))

        def work():
            if buyer is None:
                job.update(item, status="error", note="This buyer is no longer in the list.")
                return
            draft = ai.write_email(seller, buyers.buyer_text(buyer), buyers.sender_text(sender), language, anonymous)
            email = buyers.finish_email(draft, buyer, sender, language, anonymous, p)
            buyers.update_campaign(out_dir, bid, lambda c: c.setdefault("emails", {}).__setitem__(buyer["business_id"], email))
            job.update(item, status="done")
        _run_safely(job, item, work)

    with ThreadPoolExecutor(3) as pool:
        list(pool.map(one, job.items))
    job.finish()


def run_followup_job(job: Job, out_dir: Path) -> None:
    """Write one follow-up per buyer whose first email got no answer."""
    bid = job.options["seller"]
    p = _load_profile(out_dir, bid)
    c = buyers.load_campaign(out_dir, bid)
    by_id = {b["business_id"]: b for b in c.get("buyers") or []}
    emails = c.get("emails") or {}
    sender = c.get("sender") or {}
    seller = buyers.seller_text(p, c.get("brief"), c.get("teaser") or "")

    def one(item):
        buyer, first = by_id.get(item["query"]), emails.get(item["query"])
        job.update(item, status="running", name=(buyer or {}).get("name"))

        def work():
            if buyer is None or first is None:
                job.update(item, status="error", note="There is no first email to follow up.")
                return
            language, anonymous = first.get("language", "en"), first.get("anonymous", True)
            draft = ai.write_follow_up(seller, buyers.buyer_text(buyer), buyers.sender_text(sender),
                                       f"Subject: {first.get('subject')}\n\n{first.get('body')}", language, anonymous)
            follow = buyers.finish_follow_up(draft, first, sender, language, anonymous, p)

            def save(c):
                if buyer["business_id"] in (c.get("emails") or {}):
                    c["emails"][buyer["business_id"]]["follow_up"] = follow
            buyers.update_campaign(out_dir, bid, save)
            job.update(item, status="done")
        _run_safely(job, item, work)

    with ThreadPoolExecutor(3) as pool:
        list(pool.map(one, job.items))
    job.finish()


def run_job(job: Job, out_dir: Path) -> None:
    opts = job.options
    seen: dict[str, int] = {}  # business_id -> row number where it was first found
    for item in job.items:
        if job.cancel.is_set():
            job.update(item, status="skipped", note="Stopped before this row.")
            continue
        job.update(item, status="running")
        token = cs.LOG_HANDLER.set(lambda msg, item=item: job.add_log(item, msg))
        try:
            p = cs.build_profile(item["query"], website=opts.get("website") or None,
                                 with_website=opts["with_website"], with_financials=opts["with_financials"],
                                 skip_ids=set(seen))
            if p.get("duplicate"):
                job.update(item, status="duplicate", business_id=p["business_id"], name=p["name"],
                           duplicate_of_row=seen.get(p["business_id"]), note=p["note"])
            elif p.get("found"):
                with SAVE_LOCK:
                    cs.save_profile(p, out_dir)
                buyers.sync_from_profile(out_dir, p)  # a researched buyer updates the campaigns it's in
                bid = p["registry"]["business_id"]
                seen[bid] = item["row_no"]
                job.update(item, business_id=bid, name=p["registry"]["name"])
                ai_note = {}
                if opts.get("with_ai"):
                    cs.log("→ Analyzing with Claude…")
                    try:
                        analyze_profile(p, out_dir)
                    except ai.AIError as exc:
                        ai_note = {"ai_error": str(exc)}
                job.update(item, status="found", row=flatten_profile(p), **ai_note)
            else:
                job.update(item, status="not_found", note=p.get("note"))
        except RuntimeError as exc:  # PRH lookup failed
            job.update(item, status="error", note=str(exc))
        except Exception as exc:
            traceback.print_exc()
            job.update(item, status="error", note=f"Unexpected error: {exc.__class__.__name__}: {exc}")
        finally:
            cs.LOG_HANDLER.reset(token)
    job.finish()


# --------------------------------------------------------------------------
# Web app
# --------------------------------------------------------------------------

def create_app(out_dir: Path) -> Flask:
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    app = Flask(__name__, static_folder=str(STATIC), static_url_path="/static")
    app.config["SEND_FILE_MAX_AGE_DEFAULT"] = 0  # always serve the current UI files
    app.json.sort_keys = False
    jobs: dict[str, Job] = {}
    jobs_lock = threading.Lock()

    def load_profiles(ids: list[str] | None = None) -> list[dict]:
        files = ([out_dir / f"{i}.json" for i in ids] if ids is not None else
                 [f for f in out_dir.glob("*.json") if cs.BUSINESS_ID_RE.match(f.stem)])
        profiles = []
        for f in files:
            try:
                p = json.loads(f.read_text(encoding="utf-8"))
            except (OSError, ValueError):
                continue
            if p.get("found") and p.get("registry"):
                profiles.append(p)
        return profiles

    @app.before_request
    def local_only():
        # Answer only requests addressed to this computer (blocks DNS rebinding), and accept changes
        # only as JSON from the app's own pages, so other websites can't start scrapes or paid AI calls.
        host = request.host.split("]")[0] + "]" if request.host.startswith("[") else request.host.rsplit(":", 1)[0]
        if host not in ("127.0.0.1", "localhost", "[::1]"):
            abort(403)
        if request.method == "POST":
            origin = request.headers.get("Origin")
            if request.mimetype != "application/json" or (origin and urlparse(origin).netloc != request.host):
                abort(403)

    def launch(job: Job, target) -> Job:
        with jobs_lock:
            jobs[job.id] = job
            for old in list(jobs)[:-MAX_JOBS_KEPT]:
                if jobs[old].status != "running":
                    del jobs[old]
        threading.Thread(target=target, args=(job, out_dir), daemon=True).start()
        return job

    def saved_company(bid: str) -> str:
        if not cs.BUSINESS_ID_RE.match(bid) or not (out_dir / f"{bid}.json").exists():
            abort(404)
        return bid

    def requested_ids() -> list[str] | None:
        raw = request.args.get("ids")
        if raw is None:
            return None
        return [i for i in dict.fromkeys(raw.split(",")) if cs.BUSINESS_ID_RE.match(i)]

    @app.get("/")
    def index():
        return send_from_directory(STATIC, "index.html")

    @app.get("/api/ping")
    def ping():
        return jsonify(app=APP_ID)

    @app.get("/api/companies")
    def companies():
        deals = {c["business_id"]: c for c in buyers.list_campaigns(out_dir)}
        as_buyer = buyers.buyer_index(out_dir)
        profiles = load_profiles()
        names = {p["registry"]["business_id"]: p["registry"].get("name") for p in profiles}
        rows = []
        for p in profiles:
            row = flatten_profile(p)
            deal = deals.get(row["business_id"]) or {}
            analysis = p.get("ai_analysis") or {}
            row.update({
                "has_analysis": bool(analysis),
                "analysis_stale": bool(analysis) and analysis.get("based_on") != p.get("scraped_at"),
                "has_brief": bool(deal.get("has_brief")),
                "buyers": deal.get("buyers", 0), "selected": deal.get("selected", 0),
                "emails": deal.get("emails", 0), "sent": deal.get("sent", 0),
                "interested": deal.get("interested", 0), "declined": deal.get("declined", 0),
                "follow_ups_due": deal.get("follow_ups_due", 0),
                # the company is a candidate buyer for these sellers
                "buyer_for": [{**x, "name": names.get(x["seller"])} for x in as_buyer.get(row["business_id"], [])
                              if x["seller"] in names],
            })
            # where the company is in the workflow: research -> analysis -> buyers -> outreach
            row["stage"] = ("outreach" if row["emails"] else "buyers" if row["buyers"]
                            else "analyzed" if row["has_analysis"] else "new")
            rows.append(row)
        rows.sort(key=lambda r: r.get("scraped_at") or "", reverse=True)
        return jsonify(rows)

    @app.get("/api/jobs")
    def running_jobs():
        with jobs_lock:
            active = [j for j in jobs.values() if j.status == "running" and j.kind != "ask"]
        return jsonify([j.to_json() for j in sorted(active, key=lambda j: j.started)])

    def company_name(bid: str) -> str:
        try:
            return _load_profile(out_dir, bid)["registry"].get("name") or bid
        except (OSError, ValueError, KeyError):
            return bid

    @app.get("/api/companies/<bid>")
    def company(bid: str):
        saved_company(bid)
        return send_from_directory(out_dir.resolve(), f"{bid}.json", mimetype="application/json",
                                   as_attachment=request.args.get("download") == "1")

    @app.get("/api/export")
    def export():
        ids = requested_ids()
        profiles = load_profiles(ids)
        if ids is not None:  # keep the order that was asked for
            order = {bid: n for n, bid in enumerate(ids)}
            profiles.sort(key=lambda p: order[p["registry"]["business_id"]])
        else:
            profiles.sort(key=lambda p: (p["registry"].get("name") or "").lower())
        fmt = request.args.get("format", "xlsx")
        stamp = datetime.now().strftime("%Y%m%d-%H%M")
        name = (f"{profiles[0]['registry']['business_id']}" if ids and len(profiles) == 1
                else f"companies-{stamp}")
        if fmt == "json":
            body = json.dumps(profiles, ensure_ascii=False, indent=2).encode("utf-8")
            mime, ext = "application/json", "json"
        elif fmt == "csv":
            body, mime, ext = to_csv([flatten_profile(p) for p in profiles]), "text/csv; charset=utf-8", "csv"
        elif fmt == "xlsx":
            body = to_xlsx([flatten_profile(p) for p in profiles],
                           [row for p in profiles for row in evidence_rows(p)])
            mime, ext = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "xlsx"
        else:
            abort(400)
        return Response(body, mimetype=mime,
                        headers={"Content-Disposition": f'attachment; filename="{name}.{ext}"'})

    @app.post("/api/jobs")
    def start_job():
        data = request.get_json(silent=True) or {}
        queries = [str(q).strip()[:200] for q in data.get("queries") or [] if str(q).strip()]
        if not queries:
            return jsonify(error="Enter a company name or Business ID."), 400
        if len(queries) > MAX_BATCH:
            return jsonify(error=f"A batch can have at most {MAX_BATCH} companies."), 400
        kind = "lookup" if data.get("kind") == "lookup" and len(queries) == 1 else "batch"
        website = str(data.get("website") or "").strip()[:500] if kind == "lookup" else ""
        job = Job(kind, queries, {
            "website": website,
            "with_website": bool(data.get("with_website", True)),
            "with_financials": bool(data.get("with_financials", True)),
            "with_ai": kind == "batch" and bool(data.get("with_ai")) and ai.is_configured(),
        }, {"title": str(data.get("title") or "")[:120]
            or (f"Adding {queries[0]}" if kind == "lookup" else f"Adding {len(queries)} companies"),
            **({"bid": data["seller"]} if cs.BUSINESS_ID_RE.match(str(data.get("seller") or "")) else {})})
        return jsonify(launch(job, run_job).to_json()), 201

    @app.get("/api/ai/status")
    def ai_status():
        return jsonify(configured=ai.is_configured(), model=ai.MODEL)

    @app.post("/api/companies/<bid>/analysis")
    def start_analysis(bid: str):
        saved_company(bid)
        if not ai.is_configured():
            return jsonify(error="Add ANTHROPIC_API_KEY to the .env file and restart the app."), 400
        with jobs_lock:  # one analysis per company at a time
            running = next((j for j in jobs.values() if j.kind == "analysis" and j.status == "running"
                            and j.items[0]["query"] == bid), None)
        if running:
            return jsonify(running.to_json())
        detail = "detailed" if (request.get_json(silent=True) or {}).get("detail") == "detailed" else "brief"
        name = company_name(bid)
        job = Job("analysis", [bid], {"detail": detail}, {"title": f"Analyzing {name}", "bid": bid, "name": name})
        return jsonify(launch(job, run_analysis_job).to_json()), 201

    @app.post("/api/companies/<bid>/ask")
    def start_ask(bid: str):
        saved_company(bid)
        if not ai.is_configured():
            return jsonify(error="Add ANTHROPIC_API_KEY to the .env file and restart the app."), 400
        data = request.get_json(silent=True) or {}
        question = str(data.get("question") or "").strip()[:1000]
        if not question:
            return jsonify(error="Type a question."), 400
        history = [{"q": str(t["q"])[:1000], "a": str(t.get("a") or "")[:8000]}
                   for t in (data.get("history") or [])[-ai.MAX_HISTORY:] if isinstance(t, dict) and t.get("q")]
        job = Job("ask", [question], {"business_id": bid, "history": history})
        return jsonify(launch(job, run_ask_job).to_json()), 201

    # ---- selling a company: seller story, buyers, outreach emails

    def need_ai():
        if not ai.is_configured():
            return jsonify(error="Add ANTHROPIC_API_KEY to the .env file and restart the app."), 400
        return None

    def running_job(kind: str, bid: str) -> Job | None:
        with jobs_lock:
            return next((j for j in jobs.values() if j.kind == kind and j.status == "running"
                         and j.options.get("seller") == bid), None)

    @app.get("/api/campaigns")
    def campaigns():
        rows = buyers.list_campaigns(out_dir)
        names = {p["registry"]["business_id"]: p["registry"].get("name") for p in load_profiles()}
        return jsonify([{**r, "name": names.get(r["business_id"])} for r in rows if r["business_id"] in names])

    @app.get("/api/campaigns/<bid>")
    def campaign(bid: str):
        saved_company(bid)
        seller = flatten_profile(_load_profile(out_dir, bid))
        return jsonify(campaign=buyers.load_campaign(out_dir, bid), seller=seller)

    def start_campaign_job(kind: str, bid: str, target, queries: list[str], **options):
        saved_company(bid)
        error = need_ai()
        if error:
            return error
        running = running_job(kind, bid)
        if running:
            return jsonify(running.to_json())
        name = company_name(bid)
        title = {"brief": "Writing the seller story for {name}", "find": "Finding buyers for {name}",
                 "add-buyer": "Adding a buyer for {name}", "emails": "Writing {n} emails for {name}",
                 "followup": "Writing {n} follow-ups for {name}"}[kind]
        meta = {"title": title.format(name=name, n=len(queries)), "bid": bid, "name": name}
        return jsonify(launch(Job(kind, queries, {"seller": bid, **options}, meta), target).to_json()), 201

    @app.post("/api/campaigns/<bid>/brief")
    def campaign_brief(bid: str):
        return start_campaign_job("brief", bid, run_brief_job, [bid])

    @app.post("/api/campaigns/<bid>/find")
    def campaign_find(bid: str):
        return start_campaign_job("find", bid, run_find_job, [bid])

    @app.post("/api/campaigns/<bid>/buyers")
    def campaign_add_buyer(bid: str):
        query = str((request.get_json(silent=True) or {}).get("query") or "").strip()[:200]
        if not query:
            return jsonify(error="Enter a company name or Business ID."), 400
        return start_campaign_job("add-buyer", bid, run_add_buyer_job, [query])

    @app.post("/api/campaigns/<bid>/emails")
    def campaign_emails(bid: str):
        saved_company(bid)
        wanted = [str(x) for x in (request.get_json(silent=True) or {}).get("buyer_ids") or []]
        known = {b["business_id"] for b in buyers.load_campaign(out_dir, bid).get("buyers") or []}
        ids = [x for x in dict.fromkeys(wanted) if x in known][:50]
        if not ids:
            return jsonify(error="Select at least one buyer."), 400
        return start_campaign_job("emails", bid, run_emails_job, ids)

    @app.post("/api/campaigns/<bid>/followup")
    def campaign_followup(bid: str):
        saved_company(bid)
        wanted = [str(x) for x in (request.get_json(silent=True) or {}).get("buyer_ids") or []]
        sent = {k for k, e in (buyers.load_campaign(out_dir, bid).get("emails") or {}).items() if e.get("status") != "draft"}
        ids = [x for x in dict.fromkeys(wanted) if x in sent][:50]
        if not ids:
            return jsonify(error="Only emails marked as sent can get a follow-up."), 400
        return start_campaign_job("followup", bid, run_followup_job, ids)

    @app.post("/api/campaigns/<bid>/edit")
    def campaign_edit(bid: str):
        saved_company(bid)
        data = request.get_json(silent=True) or {}
        seller = _load_profile(out_dir, bid)

        def change(c):
            if "teaser" in data:
                c["teaser"] = str(data["teaser"])[:3000]
            if isinstance(data.get("sender"), dict):
                c["sender"] = {k: str(data["sender"].get(k) or "")[:200] for k in buyers.SENDER_FIELDS}
            if data.get("language") in ai.LANGUAGES:
                c["language"] = data["language"]
            if "anonymous" in data:
                c["anonymous"] = bool(data["anonymous"])
            for b in c.get("buyers") or []:
                if b["business_id"] in (data.get("selected") or {}):
                    b["selected"] = bool(data["selected"][b["business_id"]])
            if data.get("remove_buyer"):
                c["buyers"] = [b for b in c.get("buyers") or [] if b["business_id"] != data["remove_buyer"]]
                (c.get("emails") or {}).pop(data["remove_buyer"], None)
            e = data.get("email")
            if isinstance(e, dict) and e.get("buyer") in (c.get("emails") or {}):
                email = c["emails"][e["buyer"]]
                for key in ("to", "subject", "body"):
                    if key in e:
                        email[key] = str(e[key])[:20000]
                if e.get("status"):
                    buyers.set_status(email, e["status"])
                email["warnings"] = buyers.check_email(email, seller, email.get("anonymous", True))
            f = data.get("follow_up")
            if isinstance(f, dict) and ((c.get("emails") or {}).get(f.get("buyer")) or {}).get("follow_up"):
                follow = c["emails"][f["buyer"]]["follow_up"]
                for key in ("to", "subject", "body"):
                    if key in f:
                        follow[key] = str(f[key])[:20000]
                if f.get("status") in ("draft", "sent"):
                    follow["status"] = f["status"]
                    follow["sent_at"] = datetime.now().isoformat(timespec="seconds") if f["status"] == "sent" else None
                follow["warnings"] = buyers.check_email(follow, seller, c["emails"][f["buyer"]].get("anonymous", True))
        return jsonify(campaign=buyers.update_campaign(out_dir, bid, change))

    @app.get("/api/campaigns/<bid>/emails.csv")
    def campaign_csv(bid: str):
        saved_company(bid)
        return Response(buyers.emails_csv(buyers.load_campaign(out_dir, bid)), mimetype="text/csv; charset=utf-8",
                        headers={"Content-Disposition": f'attachment; filename="outreach-{bid}.csv"'})

    @app.get("/api/jobs/<job_id>")
    def job_status(job_id: str):
        job = jobs.get(job_id)
        if job is None:
            abort(404)
        return jsonify(job.to_json())

    @app.post("/api/jobs/<job_id>/stop")
    def stop_job(job_id: str):
        job = jobs.get(job_id)
        if job is None:
            abort(404)
        job.cancel.set()
        return jsonify(job.to_json())

    return app


# --------------------------------------------------------------------------
# Start-up
# --------------------------------------------------------------------------

def _port_in_use(port: int) -> bool:
    with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as s:
        return s.connect_ex(("127.0.0.1", port)) == 0


def _is_our_app(port: int) -> bool:
    try:
        return requests.get(f"http://127.0.0.1:{port}/api/ping", timeout=1).json().get("app") == APP_ID
    except (requests.RequestException, ValueError):
        return False


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="Start the Company Scraper web app.")
    ap.add_argument("--port", type=int, default=DEFAULT_PORT, help=f"Port (default {DEFAULT_PORT})")
    ap.add_argument("--out", default=str(HERE / "output"), help="Folder for saved profiles (default: output)")
    ap.add_argument("--no-browser", action="store_true", help="Don't open the browser automatically")
    args = ap.parse_args(argv)

    port = args.port
    if _port_in_use(port):
        if _is_our_app(port):
            url = f"http://127.0.0.1:{port}"
            print(f"Company Scraper is already running at {url}")
            if not args.no_browser:
                webbrowser.open(url)
            return 0
        port = 0  # something else has the port: let the OS pick a free one

    logging.getLogger("werkzeug").setLevel(logging.WARNING)  # no line per request
    server = make_server("127.0.0.1", port, create_app(Path(args.out)), threaded=True)
    url = f"http://127.0.0.1:{server.server_port}"
    print(f"\n  Company Scraper is running at {url}\n"
          f"  Profiles are saved in {Path(args.out).resolve()}\n"
          f"  Keep this window open while you use the app. Press Ctrl+C to stop.\n")
    if not args.no_browser:
        threading.Timer(0.8, webbrowser.open, [url]).start()
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        print("\nStopped.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
